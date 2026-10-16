import os
import threading
import glob
import json
import base64
try:
    import telebot
    from flask import Flask, render_template, request, jsonify, send_from_directory
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    import requests
except ImportError:
    print("Xatolik: Kutubxonalar yo'q. Terminalda yozing: pip install pyTelegramBotAPI flask openpyxl requests pillow imageio imageio-ffmpeg")
    exit(1)

from video_generator import render_dynamic_video, extract_video_from_url, try_replicate_faceswap, try_replicate_video

# XAVFSIZLIK: Kalitlarni serverdan olamiz
TELEGRAM_BOT_TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN", "")
OPENROUTER_API_KEY = os.environ.get("OPENROUTER_API_KEY", "")
OPENROUTER_MODEL = os.environ.get("OPENROUTER_MODEL", "google/gemini-2.5-flash")
if not OPENROUTER_MODEL or "llama" in OPENROUTER_MODEL.lower():
    OPENROUTER_MODEL = "google/gemini-2.5-flash"
REPLICATE_API_TOKEN = os.environ.get("REPLICATE_API_TOKEN", "")
if "REPLICATE_API_TOKEN" not in os.environ and REPLICATE_API_TOKEN:
    os.environ["REPLICATE_API_TOKEN"] = REPLICATE_API_TOKEN
OWNER_ID = os.environ.get("OWNER_ID", "")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOADS_DIR = os.path.join(BASE_DIR, 'uploads')
os.makedirs(UPLOADS_DIR, exist_ok=True)

app = Flask(
    __name__,
    template_folder=os.path.join(BASE_DIR, 'templates'),
    static_folder=os.path.join(BASE_DIR, 'static')
)

# =======================================================
# HAQIQIY EXCEL (.XLSX) YARATISH ASBOBI
# =======================================================
def excel_fayl_yaratish(fayl_nomi: str, ustunlar: list[str], qatorlar: list[list[str]]) -> str:
    """Haqiqiy professional Excel (.xlsx) fayl yaratadi."""
    if not fayl_nomi.endswith('.xlsx'):
        fayl_nomi = fayl_nomi.replace('.csv', '') + '.xlsx'
        
    fpath = os.path.join(BASE_DIR, fayl_nomi)
    print(f"[Excel Asbob]: {fayl_nomi} yaratilmoqda...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hisobot"
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    ws.append(ustunlar)
    for col_num in range(1, len(ustunlar) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        
    for row_idx, row_data in enumerate(qatorlar, start=2):
        ws.append(row_data)
        for col_num in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.font = data_font
            cell.border = thin_border
            if str(cell.value).replace(' ', '').replace(',', '').replace('.', '').isdigit():
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
    for col in ws.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_length + 4, 12)
        
    ws.row_dimensions[1].height = 26
    wb.save(fpath)
    return fayl_nomi

# =======================================================
# OPENROUTER AI ENGINE (MULTI-AGENT CORE & TOOLS)
# =======================================================
TOOLS_SCHEMA = [
    {
        "type": "function",
        "function": {
            "name": "excel_fayl_yaratish",
            "description": "Faqat foydalanuvchi hisob-kitob, smeta, narxlar yoki jadval so'raganda chaqiriladi. Professional Excel (.xlsx) jadval faylini yaratish.",
            "parameters": {
                "type": "object",
                "properties": {
                    "fayl_nomi": {
                        "type": "string",
                        "description": "Fayl nomi (masalan: hisobot_narxlar.xlsx)"
                    },
                    "ustunlar": {
                        "type": "array",
                        "items": {"type": "string"},
                        "description": "Ustunlar nomlari ro'yxati (masalan: ['T/r', 'Kategoriya', 'Vazn', 'Narx', 'Izoh'])"
                    },
                    "qatorlar": {
                        "type": "array",
                        "items": {
                            "type": "array",
                            "items": {"type": "string"}
                        },
                        "description": "Ma'lumotlar qatorlari"
                    }
                },
                "required": ["fayl_nomi", "ustunlar", "qatorlar"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "video_rolik_yaratish",
            "description": "Faqat foydalanuvchi video, rolik, klip yoki animatsiya so'raganda chaqiriladi. Haqiqiy .mp4 video rolik yaratish.",
            "parameters": {
                "type": "object",
                "properties": {
                    "video_nomi": {
                        "type": "string",
                        "description": "Video fayl nomi (masalan: reklama_rolik.mp4)"
                    },
                    "mavzu": {
                        "type": "string",
                        "description": "Video rolikning asosiy g'oyasi yoki sarlavhasi"
                    },
                    "kadrlar": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "sarlavha": {"type": "string", "description": "Kadr sarlavhasi (masalan: 1-Kadr: Kirish)"},
                                "matn": {"type": "string", "description": "Kadr davomidagi ovozli diktor nutqi yoki subtitr matni"},
                                "vizual": {"type": "string", "description": "Kadrda nimalar aks etishi"}
                            },
                            "required": ["sarlavha", "matn"]
                        },
                        "description": "Videodagi kadrma-kadr sahnalar ro'yxati"
                    }
                },
                "required": ["video_nomi", "mavzu", "kadrlar"]
            }
        }
    }
]

SYSTEM_PROMPT = (
    "Sen — nihoyatda aqlli, samimiy va professional sun'iy intellekt maslahatchisisan (Google Gemini darajasidagi yuqori intellekt).\n\n"
    "SENING MUOMALA VA ISHLASH USLUBING:\n"
    "1. FOYDALANUVCHINING GAPINI O'ZIGA ASLO QAYTARMA! 'Siz shuni aytdingiz', 'Siz shuni so'radingiz' deb gapni takrorlama. To'g'ridan-to'g'ri mohiyatga o't va o'zbek tilida sof, chiroyli, aqlli insondek gaplash.\n"
    "2. ODDIY SUHBATDA HECH QANDAY FAYL CHIQARMA:\n"
    "   - Foydalanuvchi savol bersa, maslahat so'rasa, fikr bildirar ekan — xuddi yaqin, dono do'st va maslahatchidek tushunarli, aniq va odamdek javob ber.\n"
    "   - Aslo o'zboshimchalik bilan Excel yoki video yaratma! Foydalanuvchi so'ramagan har qanday fayl ortiqcha va taqiqlangan!\n"
    "3. ANIQ TOPSHIRIQ VA ASBOBLAR ISHLASHI:\n"
    "   - Agar foydalanuvchi: 'Excel qilib ber', 'jadval tuz', 'smeta tayyorla' deb SO'RAGANDAGINA 'excel_fayl_yaratish' asbobi orqali .xlsx jadval yarat.\n"
    "   - Agar foydalanuvchi ssilka tashlab: 'prompt yozib ber', 'ssenariy tuz' desa: video va rasmlarni 1-ga-1 tahlil qilib, MiniMax (HaiLuo AI), Kling AI uchun mukammal Gollivud Super-Promptini alohida kod blokida ingliz tilida yozib ber va o'zbekcha tushuntir.\n"
    "   - Agar foydalanuvchi: 'video rolik yaratib ber' deb SO'RAGANDAGINA 'video_rolik_yaratish' asbobini ishlat.\n"
    "4. O'zbekcha xalqona iboralarni (masalan: 'oyog'ini osmondan qilib' = noto'g'ri/teskari qilib) insondek to'g'ri tushun, xato ma'noda kulgili narsa to'qima."
)

def ask_openrouter(messages, tools=None):
    """OpenRouter API orqali so'rov yuborish"""
    headers = {
        "Authorization": f"Bearer {OPENROUTER_API_KEY}",
        "Content-Type": "application/json",
        "HTTP-Referer": "https://multi-agent-bot.onrender.com",
        "X-Title": "AI Executive Suite"
    }
    
    payload = {
        "model": OPENROUTER_MODEL,
        "messages": messages,
        "max_tokens": 2500
    }
    if tools:
        payload["tools"] = tools
        payload["tool_choice"] = "auto"
        
    res = requests.post(
        "https://openrouter.ai/api/v1/chat/completions",
        headers=headers,
        json=payload,
        timeout=60
    )
    if res.status_code != 200:
        # Zaxira: Gemini 2.0 Flash ga o'tish
        if OPENROUTER_MODEL != "google/gemini-2.0-flash-001":
            print(f"[OpenRouter Ogohlantirish]: {OPENROUTER_MODEL} xato berdi ({res.status_code}). Zaxira modelga o'tilmoqda...")
            payload["model"] = "google/gemini-2.0-flash-001"
            res2 = requests.post(
                "https://openrouter.ai/api/v1/chat/completions",
                headers=headers,
                json=payload,
                timeout=60
            )
            if res2.status_code == 200:
                return res2.json()
        raise Exception(f"OpenRouter xatosi ({res.status_code}): {res.text}")
        
    return res.json()

def process_agent_request(user_text, target_agent='director', uploaded_images=None, chat_history=None):
    """Foydalanuvchi topshirig'ini bajarib, faqat so'ralgan natijani qaytaradi"""
    # 0. Ssilka tahlili: Agar matnda Instagram, TikTok, YouTube ssilka bo'lsa, haqiqiy ma'lumotni olamiz
    found_video_url = None
    video_keyframes = []
    import re
    url_match = re.search(r'(https?://[^\s]+)', user_text)
    if url_match:
        candidate_url = url_match.group(1)
        if any(d in candidate_url.lower() for d in ['instagram.com', 'tiktok.com', 'youtube.com', 'youtu.be']):
            found_video_url = candidate_url
            try:
                info = extract_video_from_url(found_video_url, download=True)
                if info.get('status') == 'success':
                    video_keyframes = info.get('keyframes_b64', [])
                    user_text += (
                        f"\n\n[SSILKADAGI VIDEO REAL MA'LUMOTLARI]:\n"
                        f"• Muallif/Kanal: {info.get('uploader')}\n"
                        f"• Video Nomi: {info.get('title')}\n"
                        f"• Video Asosiy Tavsifi: {info.get('description')}\n"
                        f"\n[QAT'IY VAZIFA]: Ushbu video syujetini 1-GA-1 AYNAN O'ZIDEK ol (xayoliy begona sportkarlar, neon shaharlar yoki oyoq bilan yozishlarni ASLO to'qima!). "
                        f"Videodagi mashinani foydalanuvchi surati va talabidagi avtomobilga (Oq rangli Chevrolet Labo, davlat raqami 01 A 457 MB), "
                        f"haydovchini esa suratdagi inson qiyofasiga almashtirgan holda MiniMax (HaiLuo AI), Kling AI uchun ingliz tilida mukammal KINEMATIK SUPER-PROMPT yoz!"
                    )
            except Exception as e:
                print("Ssilka tahlili xatosi:", e)

    # 1. Matnni tahlil qilib, foydalanuvchining asl niyatini aniqlaymiz:
    low_u = user_text.lower()
    
    # Excel so'rovi bormi?
    wants_excel = any(w in low_u for w in [
        'excel', 'jadval', 'smeta', 'hisob-kitob jadval', 'xlsx', 
        'jadval tuz', 'jadvalga sol', 'excel qilib ber', 'excelda', 
        'jadval shaklida', 'jadval ko\'rinishida', 'jadval korinishida',
        'hisobot tuz'
    ])
    
    # Prompt so'rovi (video ssilka yoki prompt/ssenariy so'ralganmi?):
    has_video_link = bool(found_video_url)
    wants_prompt = any(w in low_u for w in [
        'prompt', 'ssenariy', 'tavsif', 'super-prompt', 'super prompt'
    ]) or has_video_link
    
    # Haqiqiy video (.mp4) generatsiyasi so'ralganmi?
    has_video_word = any(w in low_u for w in ['video', 'rolik', 'klip', 'mp4', 'animatsiya'])
    has_create_word = any(w in low_u for w in ['yarat', 'tayyorla', 'qilib ber', 'yasab ber', 'generatsiya', 'render', 'chiqar'])
    wants_video_render = (has_video_word and has_create_word) and not wants_prompt

    # AGENT VA ASBOBNI ANIQ TANLASH:
    if wants_excel and not wants_video_render:
        target_agent = 'finance'
        tools_to_pass = [TOOLS_SCHEMA[0]] # FAQAT excel_fayl_yaratish
    elif wants_video_render and not wants_excel:
        target_agent = 'video'
        tools_to_pass = [TOOLS_SCHEMA[1]] # FAQAT video_rolik_yaratish
    elif wants_prompt:
        target_agent = 'video'
        tools_to_pass = None # Temur Super-Prompt yozadi, lekin og'ir fayl yaratmaydi!
    elif target_agent == 'finance':
        tools_to_pass = [TOOLS_SCHEMA[0]]
    elif target_agent == 'video':
        tools_to_pass = [TOOLS_SCHEMA[1]] if wants_video_render else None
    else:
        # ODDIY INSONIY SUHBAT, SAVOL-JAVOB, FIKR ALMASHISH
        # Asboblar berilmaydi -> HECH QANDAY EXCEL YOKI VIDEO CHIQARILMAYDI!
        tools_to_pass = None

    if target_agent == 'finance':
        system_instruction = (
            "Sen Javohir — Bosh Iqtisodchi & Moliya tahlilchisisan. Foydalanuvchi hisob-kitob, smeta yoki jadval so'radi.\n"
            "Sening vazifang: 'excel_fayl_yaratish' funksiyasini chaqirib haqiqiy chiroyli .xlsx Excel jadval yaratish! Hech qanday video qilma! "
            "Foydalanuvchining gapini o'ziga takrorlama, samimiy va professional tushuntir."
        )
    elif target_agent == 'creative':
        system_instruction = (
            "Sen Madina — Bosh Kreativ Dizaynersan. Foydalanuvchi reklama matni yoki kreativ g'oya so'radi.\n"
            "Sening vazifang: Kuchli reklama matnlari va takliflar berish. Hech qanday fayl yaratma. Samimiy va qiziqarli yoz."
        )
    elif target_agent == 'video':
        system_instruction = (
            "Sen Temur — Bosh Video Rejissyor & AI Prompt Injiniringi mutaxassisisan.\n"
            "SENING ANIQ QOIDALARING:\n"
            "1. Foydalanuvchining gapini o'ziga aslo takrorlama!\n"
            "2. Ssilkadagi video syujetini 1-GA-1 AYNAN O'ZIDEK ol (xayoliy begona sportkarlar, neon shaharlar yoki oyoq bilan yozishlarni ASLO to'qima!).\n"
            "3. Mashinani albatta foydalanuvchi talabidagi avtomobilga (Oq rangli Chevrolet Labo, davlat raqami 01 A 457 MB), haydovchini esa rasm/tavsifdagi inson qiyofasiga almashtir!\n"
            "4. HaiLuo AI (MiniMax), Kling AI va Luma uchun mukammal inglizcha KINEMATIK SUPER-PROMPTni alohida kod blokida (```) yozib ber.\n"
            "5. Foydalanuvchiga uni qaysi bepul platformalarga (HaiLuo yoki Kling) qo'yib tekinga video chiqarib olish mumkinligini o'zbek tilida londa tushuntir.\n"
            "6. Agar foydalanuvchi 'video yaratib ber' deb ochiq aytmagan bo'lsa, og'ir render asbobini chaqirma, faqat prompt va ssenariyni ber!"
        )
    else:
        system_instruction = SYSTEM_PROMPT

    # Multimodal xabar shakllantirish (AI rasmlarni va kadrlarni o'z ko'zi bilan ko'rishi uchun)
    user_msg_content = [{"type": "text", "text": user_text}]
    
    # Videodan olingan kadrlar
    for kf in video_keyframes[:2]:
        user_msg_content.append({
            "type": "image_url",
            "image_url": {"url": f"data:image/jpeg;base64,{kf}"}
        })
        
    # Foydalanuvchi yuklagan rasmlar (Labo va inson surati)
    if uploaded_images:
        from PIL import Image
        import io
        for img_p in uploaded_images[:3]:
            if os.path.exists(img_p):
                try:
                    with Image.open(img_p) as im:
                        im.thumbnail((512, 512))
                        buf = io.BytesIO()
                        im.convert('RGB').save(buf, format='JPEG', quality=80)
                        b64_im = base64.b64encode(buf.getvalue()).decode('utf-8')
                        user_msg_content.append({
                            "type": "image_url",
                            "image_url": {"url": f"data:image/jpeg;base64,{b64_im}"}
                        })
                except Exception as e:
                    print("Rasm b64 xatosi:", e)

    user_payload = user_msg_content if len(user_msg_content) > 1 else user_text

    messages = [
        {"role": "system", "content": system_instruction}
    ]
    if chat_history:
        for h in chat_history[-6:]:
            if isinstance(h, dict) and h.get('content') and h.get('role') in ['user', 'assistant']:
                messages.append({"role": h['role'], "content": str(h['content'])})
    messages.append({"role": "user", "content": user_payload})

    # So'rov yuborish
    resp = ask_openrouter(messages, tools=tools_to_pass)
    choice = resp['choices'][0]
    message_obj = choice['message']
    
    tool_calls = message_obj.get('tool_calls')
    created_excel = None
    created_video = None
    
    if tool_calls:
        tool_responses = []
        for tc in tool_calls:
            t_name = tc['function']['name']
            t_id = tc['id']
            try:
                args = json.loads(tc['function']['arguments'])
            except Exception:
                args = {}
                
            # FAQAT BITTA ASOSIY ASBOB ISHLATILADI:
            if t_name == 'excel_fayl_yaratish' and not created_video:
                try:
                    f_res = excel_fayl_yaratish(
                        args.get('fayl_nomi', 'hisobot.xlsx'),
                        args.get('ustunlar', []),
                        args.get('qatorlar', [])
                    )
                    created_excel = f_res
                    created_video = None  # Video bo'lmasligi shart!
                    content_str = f"Excel fayli '{f_res}' yaratildi."
                except Exception as e:
                    content_str = f"Excel xatosi: {e}"
                tool_responses.append({
                    "role": "tool",
                    "tool_call_id": t_id,
                    "name": t_name,
                    "content": content_str
                })
            elif t_name == 'video_rolik_yaratish' and not created_excel:
                try:
                    ref_vid = None
                    if found_video_url:
                        print(f"[Referans video]: {found_video_url} yuklab olinmoqda...")
                        v_info = extract_video_from_url(found_video_url, download=True)
                        if v_info.get('video_path') and os.path.exists(v_info['video_path']):
                            ref_vid = v_info['video_path']
                            
                    v_res = render_dynamic_video(
                        args.get('video_nomi', 'rolik.mp4'),
                        args.get('mavzu', 'Video taqdimot'),
                        args.get('kadrlar', []),
                        uploaded_images=uploaded_images,
                        reference_video=ref_vid
                    )
                    created_video = v_res
                    created_excel = None  # Excel bo'lmasligi shart!
                    content_str = f"Video fayli '{v_res}' yaratildi."
                except Exception as e:
                    content_str = f"Video generatsiya holati: {e}. Iltimos, foydalanuvchiga ssenariy asosidagi inglizcha KINEMATIK SUPER-PROMPTni alohida kod blokida va HaiLuo AI (hailuoai.video), Kling AI (klingai.com) saytlari orqali 100% bepul video olish yo'riqnomasini to'liq taqdim et."
                tool_responses.append({
                    "role": "tool",
                    "tool_call_id": t_id,
                    "name": t_name,
                    "content": content_str
                })
                
        if tool_responses:
            messages.append(message_obj)
            messages.extend(tool_responses)
            final_resp = ask_openrouter(messages)
            final_text = final_resp['choices'][0]['message'].get('content', '')
        else:
            final_text = message_obj.get('content', '')
    else:
        final_text = message_obj.get('content', '')

    # Qo'shimcha nazorat: agar foydalanuvchi aniq so'ragan bo'lsa-yu, ikkalasi aralashsa
    if created_excel and created_video:
        low_t = user_text.lower()
        if any(w in low_t for w in ['video', 'rolik', 'klip', 'ssilka', 'silka']):
            created_excel = None
        else:
            created_video = None

    if created_excel:
        ep = os.path.join(BASE_DIR, created_excel)
        if not os.path.exists(ep) or os.path.getsize(ep) < 500:
            created_excel = None

    if created_video:
        vp = os.path.join(BASE_DIR, created_video)
        if not os.path.exists(vp) or os.path.getsize(vp) < 1000:
            created_video = None

    return final_text, created_excel, created_video

# =======================================================
# FLASK WEB DASHBOARD & REST API (FON REJIMI BILAN)
# =======================================================
JOBS_STORE = {}
JOBS_LOCK = threading.Lock()

@app.errorhandler(Exception)
def handle_any_exception(e):
    """Har qanday server xatosini HTML emas, JSON formatda qaytaradi (Unexpected token '<' oldini oladi)."""
    code = getattr(e, 'code', 500)
    return jsonify({'status': 'error', 'error': str(e)}), code

@app.route('/')
def dashboard():
    return render_template('index.html')

def _background_worker(job_id, task_text, target_agent, uploaded_images):
    try:
        report, excel_file, video_file = process_agent_request(
            task_text,
            target_agent=target_agent,
            uploaded_images=uploaded_images
        )
        with JOBS_LOCK:
            JOBS_STORE[job_id] = {
                'status': 'success',
                'report': report,
                'excel_file': excel_file,
                'video_file': video_file
            }
    except Exception as e:
        with JOBS_LOCK:
            JOBS_STORE[job_id] = {
                'status': 'error',
                'error': str(e)
            }

@app.route('/api/job/<job_id>', methods=['GET'])
def api_job_status(job_id):
    with JOBS_LOCK:
        job = JOBS_STORE.get(job_id)
    if not job:
        return jsonify({'status': 'error', 'error': 'Topshiriq topilmadi'}), 404
    return jsonify(job)

@app.route('/api/task', methods=['POST'])
def api_task():
    import time
    uploaded_images = []
    
    # FormData (multipart/form-data) yoki JSON ikkalasini ham xatosiz o'qiymiz
    if request.form or request.files:
        task_text = request.form.get('task', '').strip()
        target_agent = request.form.get('target_agent', 'director')
        async_mode = request.form.get('async_mode', '') == '1'
        
        for file in request.files.getlist('files'):
            if file and file.filename:
                safe_name = f"{int(time.time())}_{os.path.basename(file.filename)}"
                save_path = os.path.join(UPLOADS_DIR, safe_name)
                file.save(save_path)
                uploaded_images.append(save_path)
    else:
        data = request.get_json(silent=True) or {}
        task_text = data.get('task', '').strip()
        target_agent = data.get('target_agent', 'director')
        async_mode = bool(data.get('async_mode', False))
        
    if not task_text and not uploaded_images:
        return jsonify({'status': 'error', 'error': 'Topshiriq matni yoki rasm kiritilmadi'}), 400
        
    if async_mode:
        job_id = os.urandom(6).hex()
        with JOBS_LOCK:
            JOBS_STORE[job_id] = {'status': 'processing'}
        t = threading.Thread(
            target=_background_worker,
            args=(job_id, task_text, target_agent, uploaded_images),
            daemon=True
        )
        t.start()
        return jsonify({'status': 'processing', 'job_id': job_id})
        
    try:
        report, excel_file, video_file = process_agent_request(
            task_text,
            target_agent=target_agent,
            uploaded_images=uploaded_images
        )
        return jsonify({
            'status': 'success',
            'report': report,
            'excel_file': excel_file,
            'video_file': video_file
        })
    except Exception as e:
        return jsonify({'status': 'error', 'error': str(e)}), 500

@app.route('/api/voice', methods=['POST'])
def api_voice():
    try:
        report, excel_file, video_file = process_agent_request(
            "Ovozli topshiriq bo'yicha to'liq hisobot tayyorlab ber."
        )
        return jsonify({
            'status': 'success',
            'report': report,
            'excel_file': excel_file,
            'video_file': video_file
        })
    except Exception as e:
        return jsonify({'status': 'error', 'error': str(e)}), 500

@app.route('/api/files', methods=['GET'])
def api_files():
    xlsx_files = [os.path.basename(f) for f in glob.glob(os.path.join(BASE_DIR, "*.xlsx")) if os.path.getsize(f) > 500]
    mp4_files = [os.path.basename(f) for f in glob.glob(os.path.join(BASE_DIR, "*.mp4")) if os.path.getsize(f) > 1000]
    return jsonify({
        'xlsx_files': xlsx_files,
        'mp4_files': mp4_files,
        'files': xlsx_files + mp4_files
    })

@app.route('/api/download/<filename>')
def api_download(filename):
    fpath = os.path.join(BASE_DIR, filename)
    if os.path.exists(fpath):
        is_video = filename.endswith('.mp4')
        return send_from_directory(BASE_DIR, filename, as_attachment=not is_video)
        
    fpath_up = os.path.join(UPLOADS_DIR, filename)
    if os.path.exists(fpath_up):
        is_video = filename.endswith('.mp4')
        return send_from_directory(UPLOADS_DIR, filename, as_attachment=not is_video)
        
    return jsonify({'status': 'error', 'error': 'Fayl topilmadi'}), 404

def run_web_server():
    port = int(os.environ.get("PORT", 8000))
    print(f"[Web Server]: Server 0.0.0.0:{port} manzilida ishga tushmoqda...")
    app.run(host="0.0.0.0", port=port, debug=False, use_reloader=False, threaded=True)

# =======================================================
# TELEGRAM BOT TIZIMI VA SUHBAT XOTIRASI
# =======================================================
CHAT_HISTORIES = {}

def get_chat_history(chat_id):
    return CHAT_HISTORIES.get(chat_id, [])

def add_to_chat_history(chat_id, user_msg, assistant_msg):
    if chat_id not in CHAT_HISTORIES:
        CHAT_HISTORIES[chat_id] = []
    CHAT_HISTORIES[chat_id].append({"role": "user", "content": user_msg})
    CHAT_HISTORIES[chat_id].append({"role": "assistant", "content": assistant_msg})
    if len(CHAT_HISTORIES[chat_id]) > 8:
        CHAT_HISTORIES[chat_id] = CHAT_HISTORIES[chat_id][-8:]

def send_long_message(bot, chat_id, text):
    """Telegramning 4096 belgilik chegarasidan oshib ketmasligi uchun xabarni xavfsiz bo'lib yuboradi"""
    if not text:
        return
    if len(text) <= 4000:
        bot.send_message(chat_id, text)
        return
    parts = []
    curr = text
    while len(curr) > 4000:
        split_idx = curr.rfind('\n', 0, 4000)
        if split_idx == -1:
            split_idx = 4000
        parts.append(curr[:split_idx])
        curr = curr[split_idx:].lstrip()
    if curr:
        parts.append(curr)
    for p in parts:
        bot.send_message(chat_id, p)

def run_telegram_bot():
    if not TELEGRAM_BOT_TOKEN:
        print("TELEGRAM_BOT_TOKEN kiritilmagan. Bot ishga tushmadi.")
        return
        
    bot = telebot.TeleBot(TELEGRAM_BOT_TOKEN)

    def check_access(message):
        if OWNER_ID and str(message.chat.id) != str(OWNER_ID):
            bot.reply_to(message, "⛔ Kechirasiz! Bu shaxsiy AI Agent boti. Undan faqat uning haqiqiy egasi foydalana oladi.")
            return False
        return True

    @bot.message_handler(commands=['start', 'myid', 'reset', 'clear'])
    def send_welcome(message):
        if not check_access(message):
            return
        CHAT_HISTORIES.pop(message.chat.id, None)
        bot.reply_to(
            message,
            f"Salom hurmatli rahbar! Men sizning Bosh AI Maslahatchi va Jamoa Rahbaringizman.\n\n"
            f"🆔 Sizning Telegram ID: `{message.chat.id}`\n\n"
            f"💬 Men bilan xuddi insondek bemalol savol-javob qilishingiz, fikr almashishingiz mumkin. "
            f"Men gaplaringizni o'zingizga takrorlamayman va so'ramagan ortiqcha fayllaringizni yaratmayman.\n\n"
            f"🎯 Qanday ishlaymiz:\n"
            f"• 💬 Oddiy suhbat / savol -> Samimiy insoniy javob va maslahat (faylsiz)\n"
            f"• 📊 'Excel qilib ber' / 'Jadval tuz' -> Moliya bo'limi tayyor .xlsx jadval yaratadi\n"
            f"• 🎬 Video ssilka / 'Prompt yoz' -> Labo (01 A 457 MB) va yuzingiz bilan 1-ga-1 Gollivud Super-Prompti beriladi\n"
            f"• 🎥 'Video yaratib ber' -> Haqiqiy video (.mp4) yaratiladi\n\n"
            f"Suhbat tarixini tozalash uchun /reset yozishingiz mumkin. Marhamat, topshirig'ingizni bering!",
            parse_mode='Markdown'
        )

    @bot.message_handler(content_types=['text'])
    def handle_text(message):
        if not check_access(message):
            return
        bot.send_chat_action(message.chat.id, 'typing')
        try:
            history = get_chat_history(message.chat.id)
            report, excel_file, video_file = process_agent_request(
                message.text,
                chat_history=history
            )
            send_long_message(bot, message.chat.id, report)
            add_to_chat_history(message.chat.id, message.text, report)
            
            # FAQAT SO'RALGAN FAYL YUBORILADI (Ikkisi birga yuborilmaydi)
            if excel_file:
                fpath = os.path.join(BASE_DIR, excel_file)
                if os.path.exists(fpath) and os.path.getsize(fpath) > 500:
                    with open(fpath, "rb") as doc:
                        bot.send_document(message.chat.id, doc, caption="📊 Siz so'ragan tayyor Excel hisoboti")
            elif video_file:
                vpath = os.path.join(BASE_DIR, video_file)
                if os.path.exists(vpath) and os.path.getsize(vpath) > 1000:
                    with open(vpath, "rb") as vid:
                        bot.send_video(message.chat.id, vid, caption="🎬 Siz so'ragan tayyor video rolik (.mp4)")
        except Exception as e:
            bot.reply_to(message, f"Xatolik yuz berdi: {e}")

    media_groups_cache = {}
    media_groups_lock = threading.Lock()

    def process_media_group(mg_id):
        with media_groups_lock:
            data = media_groups_cache.pop(mg_id, None)
        if not data:
            return
            
        chat_id = data['chat_id']
        photos = data['photos']
        caption = data['caption']
        bot.send_chat_action(chat_id, 'typing')
        
        low_c = caption.lower()
        if any(w in low_c for w in ['video', 'rolik', 'klip', 'ssilka', 'silka', 'harakatlantir']):
            target_agent = 'video'
            prompt_text = caption or "Ushbu rasmlar asosida video rolik yaratib ber."
        elif any(w in low_c for w in ['excel', 'jadval', 'narx', 'narxlari', 'smeta', 'hisob', 'hisobot']):
            target_agent = 'finance'
            prompt_text = caption or "Ushbu rasmlar bo'yicha narxlar va hisob-kitob qilib Excel jadval tayyorla."
        else:
            target_agent = 'director'
            prompt_text = caption or f"Ushbu biriktirilgan {len(photos)} ta rasmni tahlil qilib to'liq ma'lumot bering."

        try:
            history = get_chat_history(chat_id)
            report, excel_file, video_file = process_agent_request(
                prompt_text,
                target_agent=target_agent,
                uploaded_images=photos,
                chat_history=history
            )
            send_long_message(bot, chat_id, report)
            add_to_chat_history(chat_id, prompt_text, report)
            if excel_file:
                fpath = os.path.join(BASE_DIR, excel_file)
                if os.path.exists(fpath) and os.path.getsize(fpath) > 500:
                    with open(fpath, "rb") as doc:
                        bot.send_document(chat_id, doc, caption="📊 Siz so'ragan tayyor Excel hisoboti")
            elif video_file:
                vpath = os.path.join(BASE_DIR, video_file)
                if os.path.exists(vpath) and os.path.getsize(vpath) > 1000:
                    with open(vpath, "rb") as vid:
                        bot.send_video(chat_id, vid, caption="🎬 Rasmlar asosida tayyorlangan video rolik (.mp4)")
        except Exception as e:
            bot.send_message(chat_id, f"Rasmlarni tahlil qilishda xatolik: {e}")

    @bot.message_handler(content_types=['photo'])
    def handle_photo(message):
        if not check_access(message):
            return
        try:
            file_info = bot.get_file(message.photo[-1].file_id)
            downloaded = bot.download_file(file_info.file_path)
            
            photo_name = f"tg_photo_{message.photo[-1].file_id[:8]}_{os.urandom(2).hex()}.jpg"
            photo_path = os.path.join(UPLOADS_DIR, photo_name)
            with open(photo_path, 'wb') as new_f:
                new_f.write(downloaded)
                
            mg_id = getattr(message, 'media_group_id', None)
            if mg_id:
                with media_groups_lock:
                    if mg_id not in media_groups_cache:
                        media_groups_cache[mg_id] = {
                            'chat_id': message.chat.id,
                            'photos': [photo_path],
                            'caption': message.caption or '',
                            'timer': threading.Timer(1.5, process_media_group, args=[mg_id])
                        }
                        media_groups_cache[mg_id]['timer'].start()
                    else:
                        media_groups_cache[mg_id]['photos'].append(photo_path)
                        if message.caption:
                            media_groups_cache[mg_id]['caption'] = message.caption
                        media_groups_cache[mg_id]['timer'].cancel()
                        media_groups_cache[mg_id]['timer'] = threading.Timer(1.5, process_media_group, args=[mg_id])
                        media_groups_cache[mg_id]['timer'].start()
                    return

            bot.send_chat_action(message.chat.id, 'typing')
            caption = message.caption or ""
            low_c = caption.lower()
            if any(w in low_c for w in ['video', 'rolik', 'klip', 'ssilka', 'silka', 'harakatlantir']):
                target_agent = 'video'
                prompt_text = caption or "Ushbu rasm asosida video rolik yaratib ber."
            elif any(w in low_c for w in ['excel', 'jadval', 'narx', 'narxlari', 'smeta', 'hisob', 'hisobot']):
                target_agent = 'finance'
                prompt_text = caption or "Ushbu rasm bo'yicha narxlar va hisob-kitob qilib Excel jadval tayyorla."
            else:
                target_agent = 'director'
                prompt_text = caption or "Ushbu rasmni tahlil qilib to'liq ma'lumot bering."

            history = get_chat_history(message.chat.id)
            report, excel_file, video_file = process_agent_request(
                prompt_text,
                target_agent=target_agent,
                uploaded_images=[photo_path],
                chat_history=history
            )
            send_long_message(bot, message.chat.id, report)
            add_to_chat_history(message.chat.id, prompt_text, report)
            if excel_file:
                fpath = os.path.join(BASE_DIR, excel_file)
                if os.path.exists(fpath) and os.path.getsize(fpath) > 500:
                    with open(fpath, "rb") as doc:
                        bot.send_document(message.chat.id, doc, caption="📊 Siz so'ragan tayyor Excel hisoboti")
            elif video_file:
                vpath = os.path.join(BASE_DIR, video_file)
                if os.path.exists(vpath) and os.path.getsize(vpath) > 1000:
                    with open(vpath, "rb") as vid:
                        bot.send_video(message.chat.id, vid, caption="🎬 Rasm asosida tayyorlangan video rolik (.mp4)")
        except Exception as e:
            bot.reply_to(message, f"Rasm tahlilida xatolik: {e}")

    @bot.message_handler(content_types=['video'])
    def handle_video(message):
        if not check_access(message):
            return
        bot.send_chat_action(message.chat.id, 'typing')
        try:
            file_info = bot.get_file(message.video.file_id)
            downloaded = bot.download_file(file_info.file_path)
            
            vid_name = f"tg_vid_{message.video.file_id[:8]}.mp4"
            vid_path = os.path.join(UPLOADS_DIR, vid_name)
            with open(vid_path, 'wb') as new_f:
                new_f.write(downloaded)
                
            caption = message.caption or "Ushbu videoni tahlil qilib to'liq ma'lumot bering."
            history = get_chat_history(message.chat.id)
            report, excel_file, video_file = process_agent_request(
                caption,
                target_agent='video',
                uploaded_images=[vid_path],
                chat_history=history
            )
            send_long_message(bot, message.chat.id, report)
            add_to_chat_history(message.chat.id, caption, report)
            if excel_file:
                fpath = os.path.join(BASE_DIR, excel_file)
                if os.path.exists(fpath) and os.path.getsize(fpath) > 500:
                    with open(fpath, "rb") as doc:
                        bot.send_document(message.chat.id, doc, caption="📊 Hisobot (.xlsx)")
            elif video_file:
                vpath = os.path.join(BASE_DIR, video_file)
                if os.path.exists(vpath) and os.path.getsize(vpath) > 1000:
                    with open(vpath, "rb") as vid:
                        bot.send_video(message.chat.id, vid, caption="🎬 Tayyor video rolik (.mp4)")
        except Exception as e:
            bot.reply_to(message, f"Video tahlilida xatolik: {e}")

    @bot.message_handler(content_types=['voice'])
    def handle_voice(message):
        if not check_access(message):
            return
        bot.send_chat_action(message.chat.id, 'typing')
        try:
            bot.reply_to(message, "Ovozli xabar qabul qilindi. Topshiriq bajarilmoqda...")
            report, excel_file, video_file = process_agent_request(
                "Ovozli topshiriq: Iltimos, topshiriq bo'yicha to'liq hisobot bering."
            )
            bot.send_message(message.chat.id, report)
            if excel_file:
                fpath = os.path.join(BASE_DIR, excel_file)
                if os.path.exists(fpath) and os.path.getsize(fpath) > 500:
                    with open(fpath, "rb") as doc:
                        bot.send_document(message.chat.id, doc, caption="📊 Siz so'ragan tayyor Excel hisoboti")
            elif video_file:
                vpath = os.path.join(BASE_DIR, video_file)
                if os.path.exists(vpath) and os.path.getsize(vpath) > 1000:
                    with open(vpath, "rb") as vid:
                        bot.send_video(message.chat.id, vid, caption="🎬 Siz so'ragan tayyor video rolik (.mp4)")
        except Exception as e:
            bot.reply_to(message, f"Xatolik yuz berdi: {e}")

    print("Telegram bot muvaffaqiyatli ishga tushdi...")
    bot.infinity_polling(timeout=20, long_polling_timeout=20)

def start_bot_worker():
    """Telegram botni fon rejimida uzluksiz va xatosiz ishlatish uchun qo'riqchi funksiya"""
    import time
    if not TELEGRAM_BOT_TOKEN:
        print("[OGOHLANTIRISH]: TELEGRAM_BOT_TOKEN Environment Variable kiritilmagan. Telegram bot ishlamayapti.")
        return
    while True:
        try:
            print("[Telegram Bot]: Ulanmoqda...")
            run_telegram_bot()
        except Exception as e:
            print(f"[Telegram Bot Xatosi]: {e}. 5 soniyadan keyin qayta uriniladi...")
            time.sleep(5)

if __name__ == "__main__":
    t = threading.Thread(target=start_bot_worker, daemon=True)
    t.start()
    run_web_server()
